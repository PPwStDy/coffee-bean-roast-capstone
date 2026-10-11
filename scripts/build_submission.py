"""Build proposal DOCX and evidence-based PDFs with the Codex bundled document runtime."""

import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)


def read(relative, default=None):
    path = ROOT / relative
    return json.loads(path.read_text()) if path.exists() and path.stat().st_size else default


metadata = read("reports/model_metadata.json")
evaluation = metadata["test_evaluation"]
history = read("reports/training_history.json")
warm = read("reports/evidence/deployed-warm.json")
cold = read("reports/evidence/deployed-cold.json")
slo = read("monitoring/slo.json")
cost = read("reports/evidence/cost-summary.json")
registry = read("reports/evidence/azure-ml-release.json")
deployment = read("reports/evidence/candidate-deployment.json")
ci = read("reports/evidence/ci-success.json")
teardown = read("reports/evidence/teardown.json")
version = metadata["model_version"]
duration = metadata["training_duration_seconds"]
budget = cost["budget"]


def readable(text):
    prefixes = "seed|batch|epoch|version|concurrency|accuracy|loss|took|on|and|all|after|before|at|with|an|is|leaving|removes|achieved|macro-F1|min|max|are|return|plus|uses|float|RGB|unique|its|has|by|Python|LR|decay|of|were|pending|deployed|median|error|reject|remaining|historical|classified|in|the|removing|yields|restored|arXiv|than|most|was|shrank|to|smoothing|pct_start|per|Session|accept|payload|count"
    text = re.sub(rf"\b({prefixes})(?=\d)", r"\1 ", text, flags=re.IGNORECASE)
    text = re.sub(r"([,;:])(?=\d)", r"\1 ", text)
    text = re.sub(
        r"(\d)(?=epochs|minutes|images|seconds|GiB|MiB|THB|USD|vCPU|MP|px|ms|requests|copies|trials|megapixels)",
        r"\1 ",
        text,
    )
    return (
        text.replace("CC BY-SA4.0", "CC BY-SA 4.0")
        .replace("al.2022", "al. (2022)")
        .replace("Swin-T,", "Swin-T, ")
    )


proposal = [
    (
        "Problem and users",
        "We provide synchronous image-level roast classification for visual coffee inspection: Dark, Green (unroasted), Light and Medium. This revised scope uses Iteration4 Swin-T. It does not certify bean defects or SCA grades. Instructor approval of the revision remains pending.",
    ),
    (
        "Data and license",
        "Both version1 Kaggle sources by Gerry (gpiosenka) and Sakdipat Ontoum (sot2542) use CC BY-SA4.0. Deduplication removes406 copies, leaving835 training,178 validation and181 test images. Attribution, modifications and Ontoum et al.2022 Coffee Roast Intelligence are recorded in DATA_LICENSE.md.",
    ),
    (
        "Training and freshness",
        f"ImageNet initialization, seed42,20epochs and batch8 are automated and traced to Git/DVC and locked dependencies. Training took{duration / 60:.2f}minutes. Validation selects the checkpoint and serving export; the selected release achieved{evaluation['accuracy']:.1%} test accuracy on181 images. Releases update the model after validation; confidence is uncalibrated and new-camera generalization remains unmeasured.",
    ),
    (
        "Serving and performance",
        "FastAPI serves a digest-pinned ONNX image in Azure Container Apps East Asia at2vCPU/4GiB, min0/max2 replicas and concurrency2. JPEG/PNG limits are5MiB decoded and20megapixels. Freeze measured warm/cold p95/p99 plus20% after1000 requests at concurrency2 and20 verified cold trials; report image size separately.",
    ),
    (
        "Failure and reliability",
        "Black/white camera frames return422 before inference, produce JSON logs and anomaly counters, and trigger a five-minute Azure Monitor email alert. CI includes data/model/API/container checks and guard-removal regression. Candidate validation precedes traffic promotion; failed validation restores previous traffic. Actual delivery and rollback require recorded receipts.",
    ),
    (
        "Costs and ownership",
        f"Azure for Students has an800THB per-student course allowance. Conservative remaining allowance is{budget['remaining_allowance_thb']:.2f}THB; the seven-day forecast plus25% margin is{budget['forecast_thb_including_25_percent_margin']:.2f}THB. Record measured serving cost/1000 and persistent costs separately. Proposed owners: Peerawat—data/model/deployment; Yongyot—monitoring/failure/demo. Confirm actual contributions. Scope approval, clinic, presentation/Q&A and submission remain external milestones.",
    ),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.7)
for name in ["Normal", "Title", "Heading 1"]:
    doc.styles[name].font.name = "Arial"
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Normal"].paragraph_format.line_spacing = 1.02
doc.styles["Title"].font.size = Pt(20)
doc.styles["Title"].paragraph_format.space_after = Pt(6)
doc.styles["Heading 1"].font.size = Pt(11)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(5)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(2)
doc.add_paragraph("Coffee Bean Roast Classification API", "Title")
doc.add_paragraph(
    "Revised ITCS355 capstone proposal\nGroup Idea · 6688218 Peerawat · 6688230 Yongyot · 11 October 2026"
)
for heading, text in proposal:
    doc.add_paragraph(heading, "Heading 1")
    doc.add_paragraph(readable(text))
for root in (doc.element, doc.styles.element):
    for border in root.xpath(".//w:pBdr"):
        border.getparent().remove(border)
doc.save(OUT / "revised_proposal.docx")

styles = getSampleStyleSheet()
styles.add(ParagraphStyle("BodyCustom", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=9))
styles.add(ParagraphStyle("Secondary", fontName="Helvetica", fontSize=9, leading=12, spaceAfter=7))
styles.add(ParagraphStyle("TitleCustom", fontName="Helvetica-Bold", fontSize=23, leading=28, spaceAfter=18))
styles.add(
    ParagraphStyle(
        "HeadingCustom", fontName="Helvetica-Bold", fontSize=14, leading=18, spaceBefore=10, spaceAfter=8
    )
)
styles.add(ParagraphStyle("CardBody", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=7))


def p(text, style="BodyCustom"):
    return Paragraph(escape(readable(text)).replace("\n", "<br/>"), styles[style])


def table(rows, widths):
    values = [[p(str(v), "Secondary") for v in row] for row in rows]
    item = Table(values, colWidths=widths, repeatRows=1, hAlign="LEFT")
    item.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EEE9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8F7")]),
            ]
        )
    )
    return item


def footer(canvas, document):
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#52605A"))
    canvas.drawString(48, 26, "ITCS355 · Group Idea · Evidence recorded 11 October 2026")
    canvas.drawRightString(564, 26, str(document.page))


def pdf(name, content):
    SimpleDocTemplate(
        str(OUT / name), pagesize=letter, rightMargin=48, leftMargin=48, topMargin=42, bottomMargin=44
    ).build(content, onFirstPage=footer, onLaterPages=footer)


card = [p("Coffee Roast Model Card", "TitleCustom"), p(version, "HeadingCustom")]
card += [
    p("Intended use", "HeadingCustom"),
    p(
        "Educational image-level roast inspection: Dark, Green (unroasted), Light, Medium. Owner: Group Idea,6688218 Peerawat and6688230 Yongyot. This release does not certify defects, SCA grades or individual beans.",
        "CardBody",
    ),
]
card += [
    p("Data and training", "HeadingCustom"),
    p(
        f"1194 unique224px RGB images;835/178/181 train/validation/test; seed42; zero exact-pixel duplicates. ImageNet Swin-T,20epochs, batch8, AdamW and OneCycleLR. Best checkpoint: epoch{metadata['best_epoch']} by validation loss. Training:{duration / 60:.2f}minutes on local Apple M1 MPS. Original Iteration4 baseline saw all original test images; its99.75% result is historical.",
        "CardBody",
    ),
]
card += [
    p("Measured evaluation", "HeadingCustom"),
    p(
        f"Selected dynamic INT8 export:{metadata['onnx_size_mib']:.2f}MiB. Validation macro-F1:1.000 for both float and INT8; probability parity error{metadata['validation_candidates'][1]['max_probability_error']:.6f} on9 validation inputs. INT8 median{metadata['validation_candidates'][1]['latency_p50_ms']:.1f}ms versus float{metadata['validation_candidates'][0]['latency_p50_ms']:.1f}ms on this Mac. Selected release evaluated once on181 test images: accuracy{evaluation['accuracy']:.1%}, macro-F1{evaluation['macro_f1']:.3f}.",
        "CardBody",
    ),
]
card += [
    table(
        [["Test class", "Support", "Precision / recall / F1"]]
        + [
            [c, int(evaluation["classification_report"][c]["support"]), "1.000 / 1.000 / 1.000"]
            for c in metadata["classes"]
        ],
        [132, 72, 312],
    )
]
card += [
    p("Serving and limitations", "HeadingCustom"),
    p(
        "RGB224 bilinear resize and ImageNet normalization. JPEG/PNG only;5MiB and20MP limits. Mandatory SHA-256/class/preprocessing/version checks. Uniform camera frames reject422 before inference. Source/camera confounding and near duplicates remain possible; arbitrary noncoffee photos can pass the guard. Confidence is uncalibrated. New-camera performance and group-separated validation are unmeasured.",
        "CardBody",
    ),
]
card += [
    p("License and lineage", "HeadingCustom"),
    p(
        "Dataset: both version1 Kaggle sources, gpiosenka/coffee-bean-dataset-resized-224-x-224 and sot2542/coffee-bean-dataset-v1, CC BY-SA4.0. Attribution/modifications/citation: DATA_LICENSE.md. Dataset licensing does not automatically establish a weight license.",
        "CardBody",
    ),
    p(
        f"Git: {metadata['code_revision']}\nDVC: {metadata['dataset_dvc_hash']}\nONNX SHA-256: {metadata['onnx_sha256']}\nAzure ML: coffee-bean-roast-classifier version1. Full parameters/dependencies/checkpoint hash are in the immutable release metadata.",
        "Secondary",
    ),
]
pdf("model_card.pdf", card)

sections = [
    (
        "Coffee Bean Roast Capstone Report",
        [
            "This report records the operational system built for Group Idea and distinguishes successful evidence from remaining work. The project owner approved changing the supplied defect proposal to roast classification. Instructor approval of the revised scope, the Session5 clinic, live presentation/Q&A and submission remain external events.",
            "The service classifies an entire coffee image into Dark, Green, Light or Medium using Iteration4 Swin-T. Green is unroasted. Results support educational visual inspection; they do not establish individual-bean defects, roast chemistry or SCA compliance.",
            f"A clean 20-epoch run completed in {duration / 60:.2f} minutes on the 8 GiB M1. Its selected INT8 export correctly classified 181/181 test images. All 43 local tests and Linux container integration passed. The historical 99.75% baseline is preserved with a contamination warning.",
            "The ten course components are versioned code/data, automated reproducible training, registered lineage, deployed inference, CI/CD that can fail, a dashboard, a working alert, cost/1000 predictions, a one-page model card and an engineered failure. Model accuracy itself carries no marks. The repository checklist is the source of current acceptance status.",
            "Evidence files include timestamps, resource IDs, Git/DVC revisions and raw provider responses. Configuration or a generated script is not counted as proof of successful live execution. Failed attempts are preserved alongside their fixes.",
        ],
    ),
    (
        "Dataset audit and attribution",
        [
            "The supplied data has1600 readable images:1200 in the original training folder and400 in its test folder. All400 original test images have exact pixel duplicates in original training. There are1194 unique images and no conflicting labels. Removing406 copies and seed42 stratified splitting yields835/178/181.",
            "The manifest retains original names, file SHA-256, normalized pixel hashes, duplicate groups, labels and final assignments. Data contracts verify class order, counts,224px RGB integrity, manifest coverage, original hash consistency and zero exact duplicates both within and between splits. Exact-pixel deduplication does not establish independence of near duplicates or photographs of the same beans.",
            "The verified primary sources are Gerry's resized version1 dataset and Sakdipat Ontoum's version1 dataset on Kaggle, both CC BY-SA4.0. DATA_LICENSE.md records creators, links, license, the removal/resplit/renaming modifications and Ontoum et al.(2022), Coffee Roast Intelligence, arXiv2206.01841. Small test fixtures are attributed separately.",
            "Data, historical weights and the clean release use DVC pointers with a shared private Azure Blob remote. A fresh checkout restored1197 historical data/model objects and passed the full data contract; the clean release has its own pointer and upload receipt. Azure CLI/OIDC credentials stay outside source control.",
        ],
    ),
    (
        "Reproducible training and export selection",
        [
            f"The actual training source commit is{metadata['code_revision']}; dataset DVC revision is{metadata['dataset_dvc_hash']}. The dependency lock SHA-256 is{metadata['dependency_lock_sha256']}. Check out that exact commit and lock to reproduce the recorded procedure rather than silently using newer dependencies.",
            "The isolated Python3.12 environment does not inherit machine site-packages. ImageNet Swin-T initialization replaces the contaminated coffee checkpoint. Seed42 is recorded with20epochs, batch8, backbone LR5e-5, head LR1e-3, AdamW weight decay0.01, label smoothing0.05, OneCycleLR and pct_start0.2. MPS numerical behavior can differ between machines; the procedure is reproducible without promising bitwise identity.",
            f"The best validation-loss checkpoint was epoch{metadata['best_epoch']} with loss{metadata['best_validation_loss']:.6f}. No test images were evaluated during training. Local MLflow records actual parameters, dependencies, history and checkpoint. A batch32 attempt caused memory pressure while Docker was running; it was stopped before completing an epoch. Batch8 with Docker stopped completed successfully.",
            "Float32 and dynamic INT8 dense-weight exports use ONNX Runtime CPU, one thread, five warmups and all178 validation images. Selection chooses the fastest median candidate whose macro-F1 falls by no more than0.01 and whose sampled probability error is at most0.05. Both validation F1 values were1.0. INT8 was64.3ms versus float130.2ms median and shrank107.4MiB to29.6MiB.",
            "The selected INT8 release was evaluated on181 test images exactly once: accuracy1.0, macro-F11.0. A marker is written before test evaluation and repeat attempts for the same bundle are refused. Checkpoint, ONNX, TorchScript, metadata, history, evaluation, model card and hash manifest are stored together. Azure ML registry version1 preserves the semantic release version in tags and the complete metadata bundle.",
        ],
    ),
    (
        "Serving contract and release safety",
        [
            "FastAPI preserves/predict,/predict/base64,/health,/ready and/metrics. Raw JPEG/PNG and multipart file uploads are documented in OpenAPI; strict base64 JSON is supported. Returned fields include predicted_class, predicted_class_id, confidence, four probabilities, quality measurements, latency and model_version. X-Model-Version must agree with the body.",
            "Input handling checks actual format, decoded size at most5MiB and dimensions at most20megapixels before image loading. Controlled errors distinguish malformed or unreadable bytes400, unsupported image format415, oversized payload413, schema/optical rejection422 and unavailable/busy inference503. Inference runs outside the asynchronous request loop with concurrency2 and bounded waiting.",
            "Readiness checks architecture, class order, input layout, ImageNet preprocessing, mandatory model/version hashes, output shape and finite outputs. Missing, corrupted and incompatible artifacts produce failed readiness and controlled503 predictions. Liveness remains available to diagnose the process.",
            "CI restores the selected immutable bundle without retraining, checks lint and provider portability, executes data/model/API/cloud contracts, verifies export gates, mutates the optical guard in an isolated copy and runs Linux container integration. The full-training workflow is separate. Successful main CI triggers release; failed CI cannot trigger deployment.",
            "Deployment captures previous traffic, validates a digest-pinned zero-traffic candidate when an existing revision is present, executes endpoint checks and frozen latency gates, then promotes100% traffic. Candidate validation failure restores previous traffic, verifies it and deactivates the rejected candidate. Promotion/rollback evidence must reflect live provider results.",
        ],
    ),
    (
        "Azure architecture and observability",
        [
            "The approved Azure for Students subscription is in East Asia. Only rg-itcs355-6688218-capstone is owned by this capstone, with course=itcs355, student=6688218 and lab=capstone tags. Infrastructure uses Bicep validation and What-if after checking permissions, region policy, quota, provider registration and budget. The existing shared Azure ML workspace is retained.",
            "Private Blob storage provides DVC; Basic ACR stores images; a runtime managed identity pulls artifacts; GitHub OIDC provides short-lived deployment credentials. Read-only PR credentials are separate from main deployment credentials. Immutable GitHub owner/repository IDs are included in federation trust; a prior name-only trust failed and was corrected using the actual issuer/audience/subject.",
            "Container Apps uses2vCPU/4GiB, HTTP concurrency2, min0/max2, managed identity and a Linux image digest. A30second configured cooldown is recorded in Bicep; actual zero-replica checks determine cold-trial eligibility. No polling of the HTTP endpoint occurs between cold trials. Storage, registry and monitoring costs persist when serving reaches zero.",
            "Local Prometheus, Grafana and Alertmanager expose class counts, latency and camera anomaly counters. Structured JSON application logs reach Azure Log Analytics. The deployed workbook shows request status volume, server p50/p95/p99 and optical rejection events. An empty-workspace alert query required a five-minute schedule; failed provisioning attempts were retained.",
            "The engineered failure is collapsed dynamic range: black and white frames return422 before model invocation. Logs and metrics name occlusion_guard_rejection. Azure Monitor's five-minute rule sends an email through the signed-in student's action group. A fired rule or successful provider action is distinct from an actual delivered mailbox receipt; delivery evidence remains pending until verified.",
        ],
    ),
    (
        "Measured performance and operating costs",
        [
            f"Warm benchmark: {('1000 requests, concurrency2, p95 ' + str(round(warm['p95_ms'], 2)) + 'ms, p99 ' + str(round(warm['p99_ms'], 2)) + 'ms; wall time ' + str(round(warm['duration_seconds'], 2)) + 's.') if warm else 'pending deployed1000-request concurrency2 evidence.'}",
            f"Cold benchmark: {('20 trials with verified zero replicas, p95 ' + str(round(cold['p95_ms'], 2)) + 'ms, p99 ' + str(round(cold['p99_ms'], 2)) + 'ms.') if cold else 'pending20 zero-replica trials.'}",
            f"Frozen targets: {json.dumps(slo) if slo else 'pending the complete initial benchmark. Targets are measured p95/p99 plus20%, rounded upward. Later releases enforce the fixed values rather than recalibrating.'}",
            f"Conservative remaining course allowance:{budget['remaining_allowance_thb']:.2f}THB after prior subscription spending. Seven-day forecast:{budget['forecast_thb_including_25_percent_margin']:.2f}THB including25% margin. Primary East Asia retail rates and dated ECB USD/THB exchange evidence are captured. Free grants are excluded; these are estimates, not final invoices.",
            f"Serving cost per1000: {json.dumps(cost['serving_cost_per_1000']) if cost['serving_cost_per_1000'] else 'pending deployed duration. The calculator uses2vCPU/4GiB times measured wall duration at the configured maximum two replicas as a conservative upper bound, plus requests. Exact billed replica-seconds and free-grant deductions are separate.'}",
            "Storage, registry, log ingestion/alerts and training are reported separately in cost_report.md. Local MPS training incurs zero Azure compute charge; device and electricity costs are unmeasured. Budget notifications at50/80/100% of an18USD remaining-allowance budget are alerts, not a shutdown. A forecast gate blocks over-budget deployments. Scoped teardown and scheduled seven-day retention remove chargeable owned resources.",
        ],
    ),
    (
        "Acceptance evidence and academic completion",
        [
            f"Remote CI evidence: {ci.get('url') if ci else 'pending'}; result: {ci.get('conclusion') if ci else 'pending'}. Model registry: {registry.get('id') if registry else 'pending'}. Deployment promotion: {deployment.get('promoted') if deployment else 'pending'}. Teardown verified: {teardown.get('group_absent') if teardown else 'pending'}.",
            "The detailed requirement-to-evidence checklist records every acceptance scenario separately. It includes actual CI links, lineage, endpoint results, export parity, image size, warm/cold reports, dashboard captures, delivered email, cost and teardown. Historical container tests and generated templates cannot stand in for final deployed evidence.",
            "Teardown requires exact capstone group ownership and checks all resource tags. It verifies group absence and preservation of the existing shared workspace. The subscription owner separately removes only the exact capstone budget and two recorded external role assignments. Shared registry history and unrelated laboratory resources are preserved.",
            "Proposed responsibility split: Peerawat owns data/model lineage and deployment; Yongyot owns monitoring, failure demonstration and presentation. Actual contributions need team confirmation. The presentation script allocates one minute each to problem, data, architecture, live prediction, failure, regression, cost and limitations. Live timing, unexpected input and five-minute Q&A are external demonstration requirements.",
            "Instructor approval of the changed scope, Session5 architecture review, live presentation/Q&A and final repository submission remain pending. Do not mark the project fully complete until all applicable technical evidence exists and academic events are recorded. A further week should focus on independently captured images, camera/source separation, confidence calibration and broader operating-condition checks.",
            "Sources: Kaggle gpiosenka/coffee-bean-dataset-resized-224-x-224 and sot2542/coffee-bean-dataset-v1; creativecommons.org/licenses/by-sa/4.0/; arxiv.org/abs/2206.01841; learn.microsoft.com/azure/container-apps/scale-app; learn.microsoft.com/azure/container-apps/log-monitoring; prices.azure.com/api/retail/prices; ECB daily exchange rates. Course requirements: local public_teaching_mlaiops/course/project/project-brief.md and the course specification.",
        ],
    ),
]
report = []
markdown = []
for index, (title, paragraphs) in enumerate(sections):
    if index:
        report.append(PageBreak())
    report.append(p(title, "TitleCustom"))
    markdown.append("# " + title)
    for text in paragraphs:
        report.append(p(text))
        markdown.append(readable(text))
pdf("comprehensive_report.pdf", report)
(ROOT / "reports/comprehensive_report.md").write_text("\n\n".join(markdown) + "\n")
print("Created revised proposal DOCX, one-page model card PDF and comprehensive report PDF")
